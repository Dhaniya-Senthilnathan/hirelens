"""
Extracts raw text from an uploaded resume file (PDF or DOCX).

Some resumes — especially ones exported from design tools (Canva, etc.)
or scanned from paper — are "image PDFs": there is no real text layer at
all, just a picture of the page. pdfplumber (and every other text-layer
extractor) will always return empty text for these. To handle that case
we fall back to OCR (Tesseract via pytesseract) automatically whenever
the normal text extraction comes back empty or too short.
"""
import io
import os
import sys

import pdfplumber
import docx

from utils.text_cleaning import clean_text

ALLOWED_EXTENSIONS = {".pdf", ".docx"}
ALLOWED_JD_FILE_EXTENSIONS = {".pdf", ".txt"}

MIN_TEXT_LAYER_CHARS = 20  # below this, treat the PDF as having no usable text layer
OCR_DPI = 300


class UnsupportedFileTypeError(Exception):
    pass


class EmptyResumeError(Exception):
    pass


class OCRUnavailableError(Exception):
    pass


def _configure_tesseract_cmd():
    try:
        import pytesseract
    except ImportError:
        return

    env_cmd = os.environ.get("TESSERACT_CMD")
    candidates = []
    if env_cmd:
        candidates.append(env_cmd)

    if os.name == "nt" or sys.platform.startswith("win"):
        candidates.extend([
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        ])
    elif sys.platform == "darwin":
        candidates.extend([
            "/opt/homebrew/bin/tesseract",
            "/usr/local/bin/tesseract",
        ])

    for candidate in candidates:
        if candidate and os.path.exists(candidate):
            pytesseract.pytesseract.tesseract_cmd = candidate
            return


def get_extension(filename: str) -> str:
    if not filename or "." not in filename:
        return ""
    return "." + filename.rsplit(".", 1)[-1].lower()


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Text-layer extraction via pdfplumber. Returns '' for image-only PDFs."""
    text_chunks = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_chunks.append(page_text)
    return "\n".join(text_chunks)


def extract_text_from_pdf_via_ocr(file_bytes: bytes, max_pages: int = 5) -> str:
    """
    OCR fallback for scanned/image-only PDFs (no embedded text layer).
    Renders each page to an image with PyMuPDF and runs Tesseract on it.
    Capped at max_pages to keep this fast for accidental huge uploads.
    """
    try:
        import pymupdf as fitz  # PyMuPDF (new import name; falls back below)
        import pytesseract
        from PIL import Image
    except ImportError:
        try:
            import fitz  # older PyMuPDF versions only expose this name
            import pytesseract
            from PIL import Image
        except ImportError:
            # OCR dependencies not installed — caller will surface the
            # original "could not extract text" error instead.
            return ""

    _configure_tesseract_cmd()

    text_chunks = []
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception:
        return ""

    try:
        for i, page in enumerate(doc):
            if i >= max_pages:
                break
            try:
                pixmap = page.get_pixmap(dpi=OCR_DPI)
                image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                text = pytesseract.image_to_string(image)
                if text and text.strip():
                    text_chunks.append(text)
            except pytesseract.pytesseract.TesseractNotFoundError as exc:
                raise OCRUnavailableError(
                    "Tesseract OCR is not installed or not found on this machine. Please install it and restart the server (see README.md -> Prerequisites)."
                ) from exc
            except Exception:
                # Skip a page if OCR for that page fails; do not fail the whole request.
                continue
    finally:
        doc.close()

    return "\n".join(text_chunks)


def extract_text_from_docx(file_bytes: bytes) -> str:
    document = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs]

    # also pull text out of any tables in the resume
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs)


def parse_resume(filename: str, file_bytes: bytes) -> str:
    """Main entry point used by the routes layer. Returns cleaned text."""
    ext = get_extension(filename)

    if ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext or 'unknown'}'. Please upload a PDF or DOCX resume."
        )

    if ext == ".pdf":
        raw_text = extract_text_from_pdf(file_bytes)
        cleaned = clean_text(raw_text)

        if not cleaned or len(cleaned) < MIN_TEXT_LAYER_CHARS:
            # No real text layer — this is likely a scanned/flattened
            # (image-only) PDF. Fall back to OCR before giving up.
            try:
                ocr_text = extract_text_from_pdf_via_ocr(file_bytes)
            except OCRUnavailableError as exc:
                raise EmptyResumeError(
                    "Tesseract OCR is not installed or not found on this machine. Please install it and restart the server (see README.md -> Prerequisites)."
                ) from exc
            ocr_cleaned = clean_text(ocr_text)
            if ocr_cleaned and len(ocr_cleaned) >= MIN_TEXT_LAYER_CHARS:
                cleaned = ocr_cleaned
    else:
        raw_text = extract_text_from_docx(file_bytes)
        cleaned = clean_text(raw_text)

    if not cleaned or len(cleaned) < MIN_TEXT_LAYER_CHARS:
        raise EmptyResumeError(
            "Could not extract readable text from this resume, even with OCR. "
            "Try re-exporting it as a text-based PDF or a DOCX file."
        )

    return cleaned


def extract_text_from_txt(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1", errors="ignore")


def parse_job_description_file(filename: str, file_bytes: bytes) -> str:
    """Used for the optional 'upload JD as a file' flow (.txt or .pdf)."""
    ext = get_extension(filename)

    if ext not in ALLOWED_JD_FILE_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported job description file type '{ext or 'unknown'}'. Please upload a .txt or .pdf file."
        )

    if ext == ".pdf":
        raw_text = extract_text_from_pdf(file_bytes)
        cleaned = clean_text(raw_text)
        if not cleaned or len(cleaned) < 10:
            try:
                ocr_text = extract_text_from_pdf_via_ocr(file_bytes)
            except OCRUnavailableError as exc:
                raise EmptyResumeError(
                    "Tesseract OCR is not installed or not found on this machine. Please install it and restart the server (see README.md -> Prerequisites)."
                ) from exc
            ocr_cleaned = clean_text(ocr_text)
            if ocr_cleaned and len(ocr_cleaned) >= 10:
                cleaned = ocr_cleaned
    else:
        raw_text = extract_text_from_txt(file_bytes)
        cleaned = clean_text(raw_text)

    if not cleaned or len(cleaned) < 10:
        raise EmptyResumeError("Could not extract readable text from this job description file.")

    return cleaned
